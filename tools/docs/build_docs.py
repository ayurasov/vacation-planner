"""Сборка комплекта документов «Планировщика отпусков» из Markdown в PDF и DOCX.

    python3 tools/docs/build_docs.py                 # все документы
    python3 tools/docs/build_docs.py 01-architecture # один или несколько
    python3 tools/docs/build_docs.py --pdf-only ...  # без DOCX (быстрее)

Нужны pandoc 3.x, python-docx, pypdf и Node.js с зависимостями tools/docs (npm ci).
Для номеров страниц в оглавлении DOCX — LibreOffice (необязательно: без него Word
обновит оглавление при открытии).

Источники — docs/specs, docs/manuals, docs/registry, docs/performance, docs/runbooks и docs/versions.md;
схемы — docs/images (их рендерит render_diagrams.mjs). Результат — docs/pdf/<имя>.pdf и
docs/docx/<имя>.docx; комплект для реестра — в docs/pdf/registry и docs/docx/registry.
PDF собирается в два прохода: после первого номера страниц разделов читаются из ссылок
содержания и вписываются в него во втором; обложка печатается отдельно без колонтитула.
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from pypdf import PdfReader, PdfWriter

ROOT = Path(__file__).resolve().parents[2]
TOOLS = Path(__file__).resolve().parent
PDF_DIR = ROOT / "docs" / "pdf"
DOCX_DIR = ROOT / "docs" / "docx"
VERSION = json.loads((ROOT / "server" / "package.json").read_text(encoding="utf-8"))["version"]
PRODUCT = "Планировщик отпусков"
MONTHS = ["января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа",
          "сентября", "октября", "ноября", "декабря"]

# имя → (каталог, заголовок на обложке, подзаголовок)
DOCS = {
    "00-feature-audit": ("specs", "Ревизия функциональности", "Состояние as-is и целевой состав to-be"),
    "01-architecture": ("specs", "Описание архитектуры", "Платформа, компоненты, решения и план развития"),
    "02-api-spec": ("specs", "Спецификация API v1", "Протокол, ошибки, права доступа"),
    "03-backlog": ("specs", "План развития и бэклог", "Версии, эпики и истории с критериями приёмки"),
    "04-data-model": ("specs", "Модель данных", "Сущности, инварианты, справочники и миграции"),
    "05-messages": ("specs", "Тексты интерфейса", "Сообщения, ошибки и уведомления"),
    "06-ai-rules": ("specs", "Правила разработки с ИИ-ассистентами", "Процесс, ограничения и ревью"),
    "07-environments-cicd": ("specs", "Стенды, CI/CD и выпуск", "Окружения, конвейеры, версии и эксплуатация"),
    "08-tz": ("specs", "Техническое задание", "Требования к корпоративному планировщику отпусков"),
    "10-prompts": ("specs", "Промпты итераций", "Шаблоны задач для ИИ-ассистентов"),
    "11-sbom-licenses": ("specs", "SBOM и лицензии", "Состав компонентов для реестра отечественного ПО"),
    "user-guide": ("manuals", "Руководство пользователя", "Сотрудник и руководитель"),
    "admin-guide": ("manuals", "Руководство администратора", "Сотрудники, справочники, настройки и аудит"),
    "deployment-guide": ("manuals", "Руководство по развёртыванию", "Установка, обновление, резервное копирование"),
    "versions": ("", "История версий", "Выпуски, авторство и изменения"),
    # Комплект для включения в единый реестр российских программ (v1.0.0, E52-01) — в docs/pdf/registry
    "submission-package": ("registry", "Комплект документов для реестра", "Перечень и чек-лист готовности"),
    "application-data": ("registry", "Сведения о программе для заявления", "Единый реестр российских программ"),
    "functional-description": ("registry", "Описание функциональных характеристик", "Назначение, функции, сценарии и окружение"),
    "architecture": ("registry", "Описание архитектуры", "Компоненты, данные, интерфейсы и безопасность"),
    "lifecycle": ("registry", "Описание процессов жизненного цикла", "Разработка, выпуск, поддержка и персонал"),
    "expert-guide": ("registry", "Инструкция для эксперта", "Стенд, установка и проверка функций"),
    "pricing": ("registry", "Сведения о стоимости", "Стоимость и условия лицензирования"),
    "license-agreement": ("registry", "Проект лицензионного договора", "Условия использования программы"),
    "third-party-notices": ("registry", "Сторонние компоненты и лицензии", "Уведомления по лицензиям компонентов поставки"),
    # Эксплуатация и нагрузка (v1.0.0)
    "workday-model": ("performance", "Модель нагрузки", "Организация на 1000–2000 сотрудников"),
    "runbook": ("performance", "Регламент нагрузочного прогона", "Стенд, данные, доказательства и критерии"),
    "diagnostics": ("runbooks", "Диагностика замедлений и ошибок", "Симптом, метрика, причина, действие"),
}

BLUE = RGBColor(0x2F, 0x4F, 0xD1)
INK = RGBColor(0x1D, 0x24, 0x33)
MUTED = RGBColor(0x6B, 0x73, 0x85)
FONT = "Arial"


def run(cmd, **kw):
    subprocess.run(cmd, check=True, **kw)


def out_dir(base, name):
    """Комплект для реестра — отдельной папкой: его передают на экспертизу целиком."""
    return base / "registry" if DOCS[name][0] == "registry" else base


def source(name):
    folder = DOCS[name][0]
    return (ROOT / "docs" / folder / f"{name}.md") if folder else ROOT / "docs" / f"{name}.md"


def today():
    d = date.today()
    return f"{d.day} {MONTHS[d.month - 1]} {d.year}"


def footer_text(name):
    return f"{PRODUCT} {VERSION} · {DOCS[name][1]}"


# ---------------------------------------------------------------- PDF
def toc_page_numbers(pdf_path):
    """Номера страниц, на которые ведут ссылки содержания (в порядке ссылок)."""
    reader = PdfReader(str(pdf_path))
    page_index = {page.indirect_reference.idnum: i for i, page in enumerate(reader.pages)}
    named = {}
    try:
        for key, dest in reader.named_destinations.items():
            named[key] = reader.get_destination_page_number(dest)
    except Exception:  # noqa: BLE001 — разные версии Chromium пишут назначения по-разному
        pass
    numbers = []
    for page in reader.pages[1:6]:  # содержание сразу после обложки
        links = []
        for ref in page.get("/Annots") or []:
            annot = ref.get_object()
            if annot.get("/Subtype") != "/Link":
                continue
            dest = annot.get("/Dest")
            if dest is None and annot.get("/A"):
                dest = annot["/A"].get("/D")
            target = None
            if isinstance(dest, str) or hasattr(dest, "startswith"):
                target = named.get(str(dest))
            elif dest is not None:
                target = page_index.get(dest[0].idnum)
            if target is None:
                continue
            rect = annot.get("/Rect")
            links.append((-float(rect[3]), float(rect[0]), target + 1))
        if not links and numbers:
            break
        numbers.extend(n for *_pos, n in sorted(links))
    return numbers


def replace_cover(pdf, cover):
    """Первая страница — обложка без колонтитула; закладки и ссылки остальных страниц сохраняются."""
    writer = PdfWriter(clone_from=str(pdf))
    writer.insert_page(PdfReader(str(cover)).pages[0], 0)
    writer.remove_page(1)
    writer.write(str(pdf))


def pandoc_html(name, md, html):
    _folder, title, subtitle = DOCS[name]
    css = (TOOLS / "manual.css").read_text(encoding="utf-8")
    run(["pandoc", str(md), "-f", "markdown+implicit_figures+link_attributes+ascii_identifiers+pipe_tables",
         "-t", "html5", "--template", str(TOOLS / "template.html"), "--toc", "--toc-depth=2",
         "--number-sections", "-V", f"css={css}", "-V", f"cover-title={title}",
         "-V", f"subtitle={subtitle} · версия {VERSION}", "-V", f"date={today()}",
         "-M", f"title={PRODUCT} — {title}", "--resource-path", str(md.parent), "-o", str(html)])


def build_pdf(name, md):
    out_dir(PDF_DIR, name).mkdir(parents=True, exist_ok=True)
    # HTML кладётся рядом с .md, чтобы относительные пути к ../images разрешались
    html = md.with_name(f".{name}.build.html")
    pdf = out_dir(PDF_DIR, name) / f"{name}.pdf"
    script = str(TOOLS / "html_to_pdf.mjs")
    footer = footer_text(name)
    pandoc_html(name, md, html)
    try:
        run(["node", script, str(html), str(pdf), footer], cwd=TOOLS)
        numbers = toc_page_numbers(pdf)
        with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False) as fh:
            json.dump(numbers, fh)
        run(["node", script, str(html), str(pdf), footer, fh.name], cwd=TOOLS)
        Path(fh.name).unlink()
        cover = pdf.with_suffix(".cover.pdf")
        run(["node", script, str(html), str(cover), footer], cwd=TOOLS, env={**os.environ, "COVER_ONLY": "1"})
        replace_cover(pdf, cover)
        cover.unlink()
    finally:
        html.unlink(missing_ok=True)
    pages = len(PdfReader(str(pdf)).pages)
    print(f"  pdf:  {pdf.relative_to(ROOT)} ({pages} стр., пунктов содержания: {len(numbers)})")


# ---------------------------------------------------------------- DOCX
def reference_docx(path):
    data = subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"], check=True,
                          capture_output=True).stdout
    path.write_bytes(data)
    doc = Document(str(path))
    styles = {s.name: s for s in doc.styles}

    def font(style, size=None, bold=None, color=None, name=FONT):
        style.font.name = name
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), name)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            rfonts.attrib.pop(qn(attr), None)
        if size:
            style.font.size = Pt(size)
        if bold is not None:
            style.font.bold = bold
        if color is not None:
            style.font.color.rgb = color

    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in styles:
            font(styles[name], 10.5, color=INK)
            styles[name].paragraph_format.space_after = Pt(5)
            styles[name].paragraph_format.line_spacing = 1.15
    font(styles["Title"], 28, True, INK)
    font(styles["Subtitle"], 15, False, BLUE)
    if "Date" in styles:
        font(styles["Date"], 11, False, MUTED)
    for level, size in {1: 18, 2: 14, 3: 12, 4: 11}.items():
        style = styles[f"Heading {level}"]
        font(style, size, True, BLUE if level == 2 else INK)
        style.paragraph_format.space_before = Pt(14 if level == 1 else 10)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    styles["Heading 1"].paragraph_format.page_break_before = True
    for name in ("Image Caption", "Caption", "Table Caption"):
        if name in styles:
            font(styles[name], 9, False, MUTED)
            styles[name].font.italic = False
    if "Source Code" in styles:
        font(styles["Source Code"], 8.5, name="Consolas")
    if "Block Text" in styles:
        font(styles["Block Text"], 10, False, INK)
    if "TOC Heading" in styles:
        font(styles["TOC Heading"], 18, True, INK)
    doc.save(str(path))


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(table):
    el = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:color"), "D6DCEA")
        el.append(b)
    table._tbl.tblPr.append(el)


def polish_docx(path, name):
    doc = Document(str(path))
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        borders(table)
        for i, row in enumerate(table.rows):
            for cell in row.cells:
                if i == 0:
                    shade(cell, "EEF2FB")
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    for r in paragraph.runs:
                        r.font.size = Pt(9)
                        if i == 0:
                            r.font.bold = True
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Block Text":
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "F2F5FF")
            paragraph._p.get_or_add_pPr().append(shd)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = f"{footer_text(name)} · стр. "
    run_ = footer.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), kind)
            run_._r.append(fld)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = text
            run_._r.append(instr)
    for r in footer.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0x8A, 0x92, 0xA6)
    # Word обновит поля (оглавление) при открытии, если LibreOffice не заполнил их заранее
    settings = doc.settings.element
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    settings.append(upd)
    doc.save(str(path))


def build_docx(name, md, reference):
    out_dir(DOCX_DIR, name).mkdir(parents=True, exist_ok=True)
    _folder, title, subtitle = DOCS[name]
    out = out_dir(DOCX_DIR, name) / f"{name}.docx"
    run(["pandoc", str(md), "-f", "markdown+implicit_figures+link_attributes+pipe_tables", "-t", "docx",
         "--reference-doc", str(reference), "--toc", "--toc-depth=2", "--number-sections",
         "-M", "toc-title=Содержание", "-M", f"title={title}",
         "-M", f"subtitle={PRODUCT} {VERSION} · {subtitle}", "-M", f"date={today()}",
         "--resource-path", str(md.parent), "-o", str(out)])
    polish_docx(out, name)
    print(f"  docx: {out.relative_to(ROOT)}")
    return out


def main(argv):
    pdf_only = "--pdf-only" in argv
    names = [a for a in argv if not a.startswith("--")] or [n for n in DOCS if source(n).exists()]
    unknown = [n for n in names if n not in DOCS]
    if unknown:
        sys.exit(f"Неизвестные документы: {', '.join(unknown)}. Доступны: {', '.join(DOCS)}")
    docx_files = []
    with tempfile.TemporaryDirectory() as tmp:
        reference = Path(tmp) / "reference.docx"
        reference_docx(reference)
        for name in names:
            md = source(name)
            print(name)
            if not pdf_only:
                docx_files.append(build_docx(name, md, reference))
            build_pdf(name, md)
    if docx_files and shutil.which("soffice"):
        try:
            run([sys.executable, str(TOOLS / "update_docx_toc.py"), *map(str, docx_files)])
        except subprocess.CalledProcessError:
            print("  (LibreOffice не обновил оглавления — Word обновит их при открытии)")


if __name__ == "__main__":
    main(sys.argv[1:])
